"""Renderizadores locais: recebem a mesma ReportTable usada pelos formatos
Google e geram o arquivo office (docx/xlsx/pptx) em memoria, sem nenhuma
chamada de rede/Drive/credencial -- funciona mesmo sem conta Google."""
import io

from docx import Document
from docx.shared import RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.util import Inches, Pt

from services.google_report_renderers import MARCADOR_TABELA, ReportTable
from services.report_templates import ReportTemplate, default_template


def render_as_docx(table: ReportTable, template: ReportTemplate | None = None) -> bytes:
    template = template or default_template()
    document = Document()
    heading = document.add_heading(table.title, level=1)
    if template.accent_hex:
        for run in heading.runs:
            run.font.color.rgb = RGBColor.from_string(template.accent_hex)
    cols = max(len(table.headers), 1)
    doc_table = document.add_table(rows=1, cols=cols)
    doc_table.style = 'Table Grid'
    for col_index, header in enumerate(table.headers):
        cell = doc_table.rows[0].cells[col_index]
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                if template.accent_hex:
                    run.font.color.rgb = RGBColor.from_string(template.accent_hex)
    for row in table.rows:
        cells = doc_table.add_row().cells
        for col_index, value in enumerate(row):
            cells[col_index].text = str(value) if value is not None else ''
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_as_docx_from_template(template_bytes: bytes, table: ReportTable) -> bytes:
    """Abre o .docx que JA e da empresa (ela que subiu, em Relatorios >
    Templates), troca o marcador MARCADOR_TABELA por uma tabela de verdade
    com os dados -- preserva 100% da formatacao/marca da empresa, a gente
    so entra com a tabela."""
    document = Document(io.BytesIO(template_bytes))
    marcador_paragrafo = next((p for p in document.paragraphs if MARCADOR_TABELA in p.text), None)
    if marcador_paragrafo is None:
        raise ValueError(f'O template nao tem o marcador {MARCADOR_TABELA} em nenhum paragrafo.')

    cols = max(len(table.headers), 1)
    doc_table = document.add_table(rows=1, cols=cols)
    doc_table.style = 'Table Grid'
    for col_index, header in enumerate(table.headers):
        doc_table.rows[0].cells[col_index].text = header
    for row in table.rows:
        cells = doc_table.add_row().cells
        for col_index, value in enumerate(row):
            cells[col_index].text = str(value) if value is not None else ''

    # python-docx so sabe ANEXAR tabela no fim do documento (add_table) --
    # pra ela aparecer no lugar do marcador, move-se o elemento XML da
    # tabela pra logo depois do paragrafo do marcador, e aí apaga o paragrafo.
    marcador_paragrafo._p.addnext(doc_table._tbl)
    marcador_paragrafo._p.getparent().remove(marcador_paragrafo._p)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_as_xlsx(table: ReportTable, template: ReportTemplate | None = None) -> bytes:
    template = template or default_template()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = table.title[:31] or 'Relatorio'
    worksheet.append(table.headers)
    header_font = Font(bold=True, color='FFFFFF') if template.accent_hex else Font(bold=True)
    for cell in worksheet[1]:
        cell.font = header_font
        if template.accent_hex:
            cell.fill = PatternFill(start_color=template.accent_hex, end_color=template.accent_hex, fill_type='solid')
    for row in table.rows:
        worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _localizar_marcador_xlsx(workbook):
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value == MARCADOR_TABELA:
                    return worksheet, cell.row, cell.column
    return None


def render_as_xlsx_from_template(template_bytes: bytes, table: ReportTable) -> bytes:
    """Abre a .xlsx que JA e da empresa, escreve a tabela a partir da celula
    que tem o marcador MARCADOR_TABELA (o cabecalho entra bem em cima dele,
    sobrescrevendo-o) -- preserva 100% da formatacao/marca da empresa."""
    workbook = load_workbook(io.BytesIO(template_bytes))
    localizacao = _localizar_marcador_xlsx(workbook)
    if localizacao is None:
        raise ValueError(f'O template nao tem o marcador {MARCADOR_TABELA} em nenhuma celula.')
    worksheet, linha, coluna = localizacao

    valores = [table.headers, *table.rows]
    for offset_linha, linha_valores in enumerate(valores):
        for offset_coluna, valor in enumerate(linha_valores):
            cell = worksheet.cell(row=linha + offset_linha, column=coluna + offset_coluna, value=valor)
            if offset_linha == 0:
                cell.font = Font(bold=True)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _localizar_marcador_pptx(presentation):
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and MARCADOR_TABELA in shape.text_frame.text:
                return slide, shape
    return None


def render_as_pptx_from_template(template_bytes: bytes, table: ReportTable) -> bytes:
    """Abre a .pptx que JA e da empresa, troca o SHAPE que tem o marcador
    MARCADOR_TABELA por uma tabela de verdade com os dados, na mesma posicao
    e tamanho do shape original -- preserva 100% do slide master/layout da
    empresa, a gente so entra com a tabela."""
    presentation = Presentation(io.BytesIO(template_bytes))
    localizacao = _localizar_marcador_pptx(presentation)
    if localizacao is None:
        raise ValueError(f'O template nao tem o marcador {MARCADOR_TABELA} em nenhum slide.')
    slide, shape = localizacao

    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    shape._element.getparent().remove(shape._element)

    rows = len(table.rows) + 1
    cols = max(len(table.headers), 1)
    graphic_frame = slide.shapes.add_table(rows, cols, left, top, width, height)
    pptx_table = graphic_frame.table
    for col_index, header in enumerate(table.headers):
        pptx_table.cell(0, col_index).text = header or ''
    for row_index, row in enumerate(table.rows, start=1):
        for col_index, value in enumerate(row):
            pptx_table.cell(row_index, col_index).text = str(value) if value is not None else ''

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def render_as_pptx(table: ReportTable, template: ReportTemplate | None = None) -> bytes:
    template = template or default_template()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = table.title
    if template.accent_hex:
        for paragraph in slide.shapes.title.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = PptxRGBColor.from_string(template.accent_hex)
    rows = len(table.rows) + 1
    cols = max(len(table.headers), 1)
    shape = slide.shapes.add_table(rows, cols, Inches(0.5), Inches(1.5), Inches(9), Inches(0.8 * rows))
    pptx_table = shape.table
    for col_index, header in enumerate(table.headers):
        header_cell = pptx_table.cell(0, col_index)
        header_cell.text = header or ''
        if template.accent_hex:
            header_cell.fill.solid()
            header_cell.fill.fore_color.rgb = PptxRGBColor.from_string(template.accent_hex)
            for paragraph in header_cell.text_frame.paragraphs:
                paragraph.font.bold = True
                paragraph.font.size = Pt(14)
                paragraph.font.color.rgb = PptxRGBColor.from_string('FFFFFF')
    for row_index, row in enumerate(table.rows, start=1):
        for col_index, value in enumerate(row):
            pptx_table.cell(row_index, col_index).text = str(value) if value is not None else ''
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
