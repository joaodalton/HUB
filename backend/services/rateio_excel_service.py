"""Gera o formulário Copel em XLSX a partir do modelo oficial fornecido."""
import io
import re
from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from services.rateio_formulario_service import montar_tabela_formulario
from services.rateio_pdf_service import aplicar_linhas_override, validar_tabela_formulario, validar_termos_adesao_obrigatorios

ASSETS_DIR = Path(__file__).resolve().parent.parent / 'assets'
XLSX_TEMPLATE_NAME = 'Formulário para cadastro de Unidades participantes do Sistema de Compensação - Geração Compartilhada (Associações).xlsx'
LINHA_INICIAL_BENEFICIARIAS = 17
LINHA_TOTAL = 41
BENEFICIARIAS_MODELO = 24


def _template_path() -> Path:
    template = ASSETS_DIR / XLSX_TEMPLATE_NAME
    if not template.exists():
        raise ValueError('Modelo XLSX do formulário Copel não encontrado em backend/assets.')
    return template


def _carregar_modelo() -> Workbook:
    workbook = load_workbook(_template_path())
    worksheet = workbook.active
    worksheet.title = 'Formulário Copel'
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = 'A17'
    return workbook


def _preencher_texto(celula, texto: str) -> None:
    celula.value = texto
    celula.font = Font(bold=False)


def gerar_formulario_excel(
    plant_id: int, responsavel_nome: str, responsavel_cpf: str, linhas_override: list[dict] | None = None
) -> bytes:
    if not responsavel_nome.strip() or not responsavel_cpf.strip():
        raise ValueError('Nome e CPF do responsavel sao obrigatorios.')

    tabela = montar_tabela_formulario(plant_id)
    aplicar_linhas_override(tabela, linhas_override)
    validar_tabela_formulario(tabela)

    validar_termos_adesao_obrigatorios(tabela, plant_id)

    workbook = _carregar_modelo()
    worksheet = workbook.active
    linhas_extras = max(0, len(tabela['linhas']) - BENEFICIARIAS_MODELO)
    if linhas_extras:
        worksheet.unmerge_cells(start_row=LINHA_TOTAL, start_column=1, end_row=LINHA_TOTAL, end_column=4)
        worksheet.insert_rows(LINHA_TOTAL, linhas_extras)
        for linha_excel in range(LINHA_TOTAL, LINHA_TOTAL + linhas_extras):
            worksheet.row_dimensions[linha_excel].height = worksheet.row_dimensions[LINHA_TOTAL - 1].height
            for coluna in range(1, 6):
                worksheet.cell(linha_excel, coluna)._style = copy(worksheet.cell(LINHA_TOTAL - 1, coluna)._style)
        worksheet.merge_cells(start_row=LINHA_TOTAL + linhas_extras, start_column=1, end_row=LINHA_TOTAL + linhas_extras, end_column=4)
    worksheet['A4'] = re.sub(r'n[ºo]\s*_+', f'nº {tabela["ucGeradora"]}', str(worksheet['A4'].value), count=1)
    worksheet['A8'] = re.sub(r'n[ºo]\s*_+', f'nº {tabela["ucAncora"]}', str(worksheet['A8'].value), count=1)

    for linha in tabela['linhas']:
        linha_excel = LINHA_INICIAL_BENEFICIARIAS + linha['ordem'] - 1
        worksheet.cell(linha_excel, 1, linha['ordem'])
        _preencher_texto(worksheet.cell(linha_excel, 2), linha['nome'] or '')
        _preencher_texto(worksheet.cell(linha_excel, 3), linha['documento'] or '')
        _preencher_texto(worksheet.cell(linha_excel, 4), linha['ucIdentificacao'] or '')
        percentual = round(float(linha['percentual']) / 100, 4)
        worksheet.cell(linha_excel, 5, percentual).number_format = '0.00%'

    total = round(sum(float(linha['percentual']) for linha in tabela['linhas']) / 100, 4)
    linha_total = LINHA_TOTAL + linhas_extras
    worksheet.cell(linha_total, 5, total).number_format = '0.00%'
    _preencher_texto(worksheet[f'C{71 + linhas_extras}'], tabela['empresaNome'] or '')
    _preencher_texto(worksheet[f'C{72 + linhas_extras}'], tabela['empresaEmail'] or '')
    _preencher_texto(worksheet[f'C{73 + linhas_extras}'], tabela['empresaCnpj'] or '')
    _preencher_texto(worksheet[f'C{75 + linhas_extras}'], responsavel_nome.strip())
    _preencher_texto(worksheet[f'C{76 + linhas_extras}'], responsavel_cpf.strip())
    worksheet[f'A{81 + linhas_extras}'] = f'Data: {date.today():%d/%m/%Y}'

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
