import io
import unittest

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Emu

from services.google_report_renderers import MARCADOR_TABELA, ReportTable
from services.office_report_renderers import (
    render_as_docx, render_as_docx_from_template, render_as_pptx, render_as_pptx_from_template,
    render_as_xlsx, render_as_xlsx_from_template,
)
from services.report_templates import SYSTEM_TEMPLATES


def _template_bytes(paragrafos: list[str]) -> bytes:
    document = Document()
    for texto in paragrafos:
        document.add_paragraph(texto)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_template_bytes(linhas: list[list]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    for linha in linhas:
        worksheet.append(linha)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pptx_template_bytes(texto_shape: str, left=Emu(100), top=Emu(200), width=Emu(300), height=Emu(400)) -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    caixa = slide.shapes.add_textbox(left, top, width, height)
    caixa.text_frame.text = texto_shape
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _table():
    return ReportTable(title='Cadastro', headers=['Nome', 'Codigo'], rows=[['Cliente A', 'UC-1'], ['Cliente B', 'UC-2']])


class OfficeReportRenderersTest(unittest.TestCase):
    def test_render_as_docx_produces_bytes_with_title_and_table(self):
        data = render_as_docx(_table())
        self.assertIsInstance(data, bytes)
        self.assertGreater(len(data), 0)

        document = Document(io.BytesIO(data))
        texts = '\n'.join(p.text for p in document.paragraphs)
        self.assertIn('Cadastro', texts)
        table = document.tables[0]
        self.assertEqual([cell.text for cell in table.rows[0].cells], ['Nome', 'Codigo'])
        self.assertEqual([cell.text for cell in table.rows[1].cells], ['Cliente A', 'UC-1'])

    def test_render_as_xlsx_produces_bytes_with_headers_and_rows(self):
        data = render_as_xlsx(_table())
        self.assertIsInstance(data, bytes)
        self.assertGreater(len(data), 0)

        workbook = load_workbook(io.BytesIO(data))
        worksheet = workbook.active
        self.assertEqual([cell.value for cell in worksheet[1]], ['Nome', 'Codigo'])
        self.assertEqual([cell.value for cell in worksheet[2]], ['Cliente A', 'UC-1'])

    def test_render_as_pptx_produces_bytes_with_title_and_table(self):
        data = render_as_pptx(_table())
        self.assertIsInstance(data, bytes)
        self.assertGreater(len(data), 0)

        presentation = Presentation(io.BytesIO(data))
        slide = presentation.slides[0]
        self.assertEqual(slide.shapes.title.text, 'Cadastro')
        table_shape = next(shape for shape in slide.shapes if shape.has_table)
        table = table_shape.table
        self.assertEqual(table.cell(0, 0).text, 'Nome')
        self.assertEqual(table.cell(1, 0).text, 'Cliente A')

    def test_empty_rows_do_not_raise(self):
        table = ReportTable(title='Vazio', headers=['Nome'], rows=[])
        self.assertTrue(render_as_docx(table))
        self.assertTrue(render_as_xlsx(table))
        self.assertTrue(render_as_pptx(table))

    def test_template_padrao_does_not_change_output(self):
        # 'padrao' = accent_hex None -- mesma saida de antes da feature de template.
        sem_template = render_as_xlsx(_table())
        com_template_padrao = render_as_xlsx(_table(), template=SYSTEM_TEMPLATES['padrao'])
        self.assertEqual(sem_template, com_template_padrao)

    def test_template_azul_aplica_cor_no_cabecalho_docx(self):
        azul = SYSTEM_TEMPLATES['azul']
        data = render_as_docx(_table(), template=azul)
        document = Document(io.BytesIO(data))
        header_run = document.tables[0].rows[0].cells[0].paragraphs[0].runs[0]
        self.assertTrue(header_run.font.bold)
        self.assertEqual(str(header_run.font.color.rgb), azul.accent_hex)

    def test_template_azul_aplica_cor_no_cabecalho_xlsx(self):
        azul = SYSTEM_TEMPLATES['azul']
        data = render_as_xlsx(_table(), template=azul)
        workbook = load_workbook(io.BytesIO(data))
        header_cell = workbook.active['A1']
        self.assertEqual(header_cell.fill.start_color.rgb[-6:], azul.accent_hex)

    def test_template_azul_aplica_cor_no_cabecalho_pptx(self):
        azul = SYSTEM_TEMPLATES['azul']
        data = render_as_pptx(_table(), template=azul)
        presentation = Presentation(io.BytesIO(data))
        table_shape = next(shape for shape in presentation.slides[0].shapes if shape.has_table)
        header_cell = table_shape.table.cell(0, 0)
        self.assertEqual(str(header_cell.fill.fore_color.rgb), azul.accent_hex)

    def test_render_as_docx_from_template_insere_tabela_no_lugar_do_marcador(self):
        template = _template_bytes(['Marca da empresa', MARCADOR_TABELA, 'Rodape'])
        data = render_as_docx_from_template(template, _table())
        document = Document(io.BytesIO(data))

        textos = [p.text for p in document.paragraphs]
        self.assertIn('Marca da empresa', textos)
        self.assertIn('Rodape', textos)
        self.assertFalse(any(MARCADOR_TABELA in texto for texto in textos))  # marcador foi removido

        self.assertEqual(len(document.tables), 1)
        table = document.tables[0]
        self.assertEqual([cell.text for cell in table.rows[0].cells], ['Nome', 'Codigo'])
        self.assertEqual([cell.text for cell in table.rows[1].cells], ['Cliente A', 'UC-1'])

        # tabela entrou FISICAMENTE entre "Marca da empresa" e "Rodape"
        # (paragrafo do marcador foi removido, so a tabela ficou no lugar dele)
        corpo = [el.tag.split('}')[-1] for el in document.element.body]
        self.assertEqual(corpo[:3], ['p', 'tbl', 'p'])

    def test_render_as_docx_from_template_sem_marcador_levanta_erro(self):
        template = _template_bytes(['Sem marcador nenhum aqui'])
        with self.assertRaises(ValueError):
            render_as_docx_from_template(template, _table())

    def test_render_as_xlsx_from_template_sobrescreve_marcador_e_preserva_resto(self):
        template = _xlsx_template_bytes([['Marca da empresa'], [], ['', MARCADOR_TABELA]])
        data = render_as_xlsx_from_template(template, _table())
        workbook = load_workbook(io.BytesIO(data))
        worksheet = workbook.active

        self.assertEqual(worksheet['A1'].value, 'Marca da empresa')  # preservado
        self.assertEqual(worksheet['B3'].value, 'Nome')  # marcador sobrescrito pelo cabecalho
        self.assertEqual(worksheet['C3'].value, 'Codigo')
        self.assertEqual(worksheet['B4'].value, 'Cliente A')
        self.assertEqual(worksheet['C4'].value, 'UC-1')
        self.assertEqual(worksheet['B5'].value, 'Cliente B')
        self.assertTrue(worksheet['B3'].font.bold)

    def test_render_as_xlsx_from_template_sem_marcador_levanta_erro(self):
        template = _xlsx_template_bytes([['Sem marcador aqui']])
        with self.assertRaises(ValueError):
            render_as_xlsx_from_template(template, _table())

    def test_render_as_pptx_from_template_troca_shape_do_marcador_por_tabela_na_mesma_posicao(self):
        template = _pptx_template_bytes(MARCADOR_TABELA, left=Emu(111), top=Emu(222), width=Emu(3330000), height=Emu(4440000))
        data = render_as_pptx_from_template(template, _table())
        presentation = Presentation(io.BytesIO(data))
        slide = presentation.slides[0]

        shapes_com_texto_marcador = [s for s in slide.shapes if s.has_text_frame and MARCADOR_TABELA in s.text_frame.text]
        self.assertEqual(shapes_com_texto_marcador, [])  # shape do marcador foi removido

        table_shape = next(s for s in slide.shapes if s.has_table)
        self.assertEqual(table_shape.left, 111)
        self.assertEqual(table_shape.top, 222)
        self.assertEqual([cell.text for cell in table_shape.table.rows[0].cells], ['Nome', 'Codigo'])
        self.assertEqual([cell.text for cell in table_shape.table.rows[1].cells], ['Cliente A', 'UC-1'])

    def test_render_as_pptx_from_template_sem_marcador_levanta_erro(self):
        template = _pptx_template_bytes('Sem marcador aqui')
        with self.assertRaises(ValueError):
            render_as_pptx_from_template(template, _table())


if __name__ == '__main__':
    unittest.main()
