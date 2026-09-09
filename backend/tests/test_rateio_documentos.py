import io
import os
import sys
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from pypdf import PdfReader

_DB = tempfile.NamedTemporaryFile(suffix='.db', delete=False)
_DB.close()
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import create_app
from extensions import db
from models.empresa import Empresa
from models.plant import Plant
from models.user import User
from services.rateio_excel_service import gerar_formulario_excel
from services.rateio_pdf_service import DriveUnavailableError, gerar_formulario_pdf
from utils.auth import generate_token
try:
    from .support import IsolatedTestRuntime
except ImportError:
    from support import IsolatedTestRuntime


def _tabela(percentual=50.0):
    return {
        'plantId': 1, 'plantNome': 'Usina A', 'ucGeradora': 'GER-1', 'ucAncora': 'GER-1',
        'empresaNome': 'Associacao A', 'empresaCnpj': '12345678000199', 'empresaEmail': 'a@test.local',
        'documentoCnpjOk': True, 'documentoEstatutoOk': True,
        'linhas': [{
            'ordem': 1, 'nome': 'Ana', 'documento': '12345678901', 'ucIdentificacao': 'UC-1',
            'percentual': percentual, 'termoAdesaoOk': True, 'clienteId': 1, 'ucId': 1,
        }],
        'somaPercentual': percentual,
    }


class RateioDocumentosTest(IsolatedTestRuntime, unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.prepare_test_runtime(f"sqlite:///{_DB.name.replace(chr(92), '/')}", 'rateio-documentos-test')
        cls.app = create_app()
        cls.app.config['TESTING'] = True
        with cls.app.app_context():
            db.create_all()
            empresa = Empresa(nome='Empresa A', slug='rateio-doc-a')
            db.session.add(empresa)
            db.session.flush()
            user = User(empresa_id=empresa.id, nome='Owner', email='owner@rateio.test', password_hash='x', role='owner')
            db.session.add(user)
            empresa_b = Empresa(nome='Empresa B', slug='rateio-doc-b')
            db.session.add(empresa_b)
            db.session.flush()
            plant_b = Plant(empresa_id=empresa_b.id, nome='Usina B', uc='UC-B', kw_pico=1)
            db.session.add(plant_b)
            db.session.commit()
            cls.user_id = user.id
            cls.plant_b_id = plant_b.id

    @classmethod
    def tearDownClass(cls):
        with cls.app.app_context():
            db.session.remove()
            db.drop_all()
            db.engine.dispose()
        os.unlink(_DB.name)
        cls.restore_test_runtime()

    def test_pdf_and_excel_reject_total_over_100_percent(self):
        tabela = _tabela(100.01)
        with patch('services.rateio_pdf_service.montar_tabela_formulario', return_value=tabela):
            with self.assertRaisesRegex(ValueError, 'excede 100'):
                gerar_formulario_pdf(1, 'Responsavel', '12345678901')
        with patch('services.rateio_excel_service.montar_tabela_formulario', return_value=tabela):
            with self.assertRaisesRegex(ValueError, 'excede 100'):
                gerar_formulario_excel(1, 'Responsavel', '12345678901')

    def test_document_outputs_expand_beyond_24_beneficiaries(self):
        tabela = _tabela(4)
        tabela['linhas'] = [dict(deepcopy(tabela['linhas'][0]), ordem=ordem, ucId=ordem) for ordem in range(1, 26)]
        with patch('services.rateio_pdf_service.montar_tabela_formulario', return_value=tabela):
            with patch('services.rateio_pdf_service.verificar_termos_adesao', return_value={'ok': True, 'faltando': []}):
                pdf = gerar_formulario_pdf(1, 'Responsavel', '12345678901')
        self.assertEqual(len(PdfReader(io.BytesIO(pdf)).pages), 5)
        with patch('services.rateio_excel_service.montar_tabela_formulario', return_value=tabela):
            with patch('services.rateio_pdf_service.verificar_termos_adesao', return_value={'ok': True, 'faltando': []}):
                workbook = load_workbook(io.BytesIO(gerar_formulario_excel(1, 'Responsavel', '12345678901')))
        sheet = workbook['Formulário Copel']
        self.assertEqual(sheet['A41'].value, 25)
        self.assertEqual(sheet['B41'].value, 'Ana')
        self.assertEqual(sheet['E42'].value, 1)
        self.assertEqual(sheet['C72'].value, 'Associacao A')

    def test_excel_is_valid_and_contains_rateio_data(self):
        with patch('services.rateio_excel_service.montar_tabela_formulario', return_value=_tabela()):
            with patch('services.rateio_pdf_service.verificar_termos_adesao', return_value={'ok': True, 'faltando': []}):
                workbook = load_workbook(io.BytesIO(gerar_formulario_excel(
                    1, 'Responsavel', '12345678901', [{'ucId': 1, 'nome': 'Ana Revisada', 'percentual': 55}]
                )))

        sheet = workbook['Formulário Copel']
        self.assertIn('GER-1', sheet['A4'].value)
        self.assertIn('GER-1', sheet['A8'].value)
        self.assertEqual(sheet['B17'].value, 'Ana Revisada')
        self.assertEqual(sheet['C17'].value, '12345678901')
        self.assertEqual(sheet['D17'].value, 'UC-1')
        self.assertEqual(sheet['E17'].value, 0.55)
        self.assertEqual(sheet['E41'].value, 0.55)
        self.assertEqual(sheet['C71'].value, 'Associacao A')
        self.assertEqual(sheet['C75'].value, 'Responsavel')

    def test_document_requirements_can_be_disabled_for_internal_testing(self):
        tabela = _tabela()
        tabela['documentoCnpjOk'] = False
        tabela['documentoEstatutoOk'] = False
        tabela['regrasDocumentos'] = {
            'documentoCnpjObrigatorio': False,
            'documentoEstatutoObrigatorio': False,
            'termosAdesaoObrigatorios': False,
        }

        with patch('services.rateio_pdf_service.montar_tabela_formulario', return_value=tabela):
            pdf = gerar_formulario_pdf(1, 'Responsavel', '12345678901')
        self.assertGreater(len(pdf), 0)

        with patch('services.rateio_excel_service.montar_tabela_formulario', return_value=tabela):
            workbook = load_workbook(io.BytesIO(gerar_formulario_excel(1, 'Responsavel', '12345678901')))
        self.assertEqual(workbook['Formulário Copel']['B17'].value, 'Ana')

    def test_rateio_document_settings_are_accepted(self):
        with self.app.app_context():
            token = generate_token(self.user_id)
        response = self.app.test_client().put(
            '/api/v1/settings',
            headers={'Authorization': f'Bearer {token}'},
            json={
                'rateioExigirDocumentoCnpj': 'false',
                'rateioExigirDocumentoEstatuto': 'false',
                'rateioExigirTermosAdesao': 'false',
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()['data']['rateioExigirTermosAdesao'], 'false')

    def test_drive_failure_on_terms_is_a_clear_503(self):
        with self.app.app_context():
            token = generate_token(self.user_id)
        with patch('routes.rateio_routes.gerar_termos_adesao_pdf', side_effect=DriveUnavailableError('OAuth ausente')):
            response = self.app.test_client().post(
                '/api/v1/rateio/formulario/gerar-termos',
                headers={'Authorization': f'Bearer {token}'}, json={'plantId': 1},
            )
        self.assertEqual(response.status_code, 503)
        self.assertIn('Google Drive nao configurado ou indisponivel', response.get_json()['error'])

    def test_pdf_generation_route_does_not_depend_on_drive_download(self):
        with self.app.app_context():
            token = generate_token(self.user_id)
        with patch('routes.rateio_routes.gerar_formulario_pdf', return_value=b'%PDF-1.4'):
            response = self.app.test_client().post(
                '/api/v1/rateio/formulario/gerar-pdf',
                headers={'Authorization': f'Bearer {token}'},
                json={'plantId': 1, 'responsavelNome': 'Responsavel', 'responsavelCpf': '12345678901'},
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, 'application/pdf')

    def test_excel_cannot_generate_a_foreign_tenant_plant(self):
        with self.app.app_context():
            token = generate_token(self.user_id)
        response = self.app.test_client().post(
            '/api/v1/rateio/formulario/gerar-excel',
            headers={'Authorization': f'Bearer {token}'},
            json={'plantId': self.plant_b_id, 'responsavelNome': 'Responsavel', 'responsavelCpf': '12345678901'},
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('Usina nao encontrada', response.get_json()['error'])


if __name__ == '__main__':
    unittest.main()
